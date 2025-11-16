"""
DocumentEncryptor单元测试
"""

import io
import pytest
from pathlib import Path
from core.engine.encryptor import DocumentEncryptor


@pytest.fixture
def encryptor():
    """创建DocumentEncryptor实例"""
    return DocumentEncryptor()


@pytest.fixture
def simple_pdf_bytes():
    """创建一个简单的PDF用于测试"""
    try:
        from PyPDF2 import PdfWriter
    except ImportError:
        pytest.skip("需要安装PyPDF2库")
    
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)  # Letter size
    
    output = io.BytesIO()
    writer.write(output)
    output.seek(0)
    return output.read()


@pytest.fixture
def simple_docx_bytes():
    """创建一个简单的Word文档用于测试"""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("需要安装python-docx库")
    
    doc = Document()
    doc.add_heading('测试文档', 0)
    doc.add_paragraph('这是一个测试文档。')
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()


class TestDocumentEncryptor:
    """DocumentEncryptor测试类"""
    
    def test_encrypt_pdf_success(self, encryptor, simple_pdf_bytes):
        """测试PDF加密成功"""
        password = "test123"
        
        # 加密PDF
        encrypted_bytes = encryptor.encrypt_pdf(simple_pdf_bytes, password)
        
        # 验证加密后的字节不为空
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0
        
        # 验证加密后的PDF可以被PyPDF2读取
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(encrypted_bytes))
            # 验证PDF已加密
            assert reader.is_encrypted
        except ImportError:
            pytest.skip("需要安装PyPDF2库验证加密结果")
    
    def test_encrypt_pdf_with_owner_password(self, encryptor, simple_pdf_bytes):
        """测试PDF使用所有者密码加密"""
        user_password = "user123"
        owner_password = "owner456"
        
        # 加密PDF
        encrypted_bytes = encryptor.encrypt_pdf(
            simple_pdf_bytes,
            user_password,
            owner_password=owner_password
        )
        
        # 验证加密成功
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0
    
    def test_encrypt_pdf_empty_password(self, encryptor, simple_pdf_bytes):
        """测试PDF加密时密码为空"""
        with pytest.raises(ValueError, match="密码不能为空"):
            encryptor.encrypt_pdf(simple_pdf_bytes, "")
    
    def test_encrypt_pdf_invalid_data(self, encryptor):
        """测试PDF加密时数据无效"""
        try:
            import PyPDF2
        except ImportError:
            pytest.skip("需要安装PyPDF2库")
        
        with pytest.raises(ValueError, match="PDF加密失败"):
            encryptor.encrypt_pdf(b"invalid pdf data", "test123")
    
    def test_encrypt_word_success(self, encryptor, simple_docx_bytes):
        """测试Word加密成功"""
        password = "test456"
        
        # 加密Word文档
        encrypted_bytes = encryptor.encrypt_word(simple_docx_bytes, password)
        
        # 验证加密后的字节不为空
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0
        
        # 验证加密后的文件大小合理（通常比原文件稍大）
        assert len(encrypted_bytes) >= len(simple_docx_bytes) * 0.8
    
    def test_encrypt_word_empty_password(self, encryptor, simple_docx_bytes):
        """测试Word加密时密码为空"""
        with pytest.raises(ValueError, match="密码不能为空"):
            encryptor.encrypt_word(simple_docx_bytes, "")
    
    def test_encrypt_word_invalid_data(self, encryptor):
        """测试Word加密时数据无效"""
        try:
            import msoffcrypto
        except ImportError:
            pytest.skip("需要安装msoffcrypto-tool库")
        
        with pytest.raises(ValueError, match="Word加密失败"):
            encryptor.encrypt_word(b"invalid docx data", "test123")
    
    def test_encrypt_document_pdf(self, encryptor, simple_pdf_bytes):
        """测试encrypt_document方法加密PDF"""
        password = "test789"
        
        # 使用通用方法加密PDF
        encrypted_bytes = encryptor.encrypt_document(
            simple_pdf_bytes,
            password,
            format="pdf"
        )
        
        # 验证加密成功
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0
    
    def test_encrypt_document_docx(self, encryptor, simple_docx_bytes):
        """测试encrypt_document方法加密Word"""
        password = "test789"
        
        # 使用通用方法加密Word
        encrypted_bytes = encryptor.encrypt_document(
            simple_docx_bytes,
            password,
            format="docx"
        )
        
        # 验证加密成功
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0
    
    def test_encrypt_document_unsupported_format(self, encryptor):
        """测试加密不支持的格式"""
        with pytest.raises(ValueError, match="不支持的文档格式"):
            encryptor.encrypt_document(
                b"some data",
                "test123",
                format="txt"
            )
    
    def test_encrypt_document_with_extra_params(self, encryptor, simple_pdf_bytes):
        """测试加密时传递额外参数"""
        password = "test123"
        
        # 使用额外参数加密PDF
        encrypted_bytes = encryptor.encrypt_document(
            simple_pdf_bytes,
            password,
            format="pdf",
            owner_password="owner456"
        )
        
        # 验证加密成功
        assert encrypted_bytes
        assert len(encrypted_bytes) > 0


if __name__ == "__main__":
    pytest.main([__file__, "-v"])

