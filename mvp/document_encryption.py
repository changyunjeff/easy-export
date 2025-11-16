#!/usr/bin/env python3
"""
文档加密MVP示例
演示如何对PDF和Word文档进行密码保护
"""

import io
import sys
from pathlib import Path

# 添加项目根目录到路径
sys.path.insert(0, str(Path(__file__).parent.parent))


def encrypt_pdf_pypdf2(input_pdf_bytes: bytes, password: str) -> bytes:
    """
    使用PyPDF2加密PDF文档
    
    Args:
        input_pdf_bytes: 输入PDF字节内容
        password: 加密密码
        
    Returns:
        加密后的PDF字节内容
    """
    try:
        from PyPDF2 import PdfReader, PdfWriter
    except ImportError:
        print("请安装PyPDF2: pip install PyPDF2")
        raise
    
    # 读取PDF
    reader = PdfReader(io.BytesIO(input_pdf_bytes))
    writer = PdfWriter()
    
    # 复制所有页面
    for page in reader.pages:
        writer.add_page(page)
    
    # 设置密码加密
    writer.encrypt(user_password=password, owner_password=None)
    
    # 保存到字节流
    output = io.BytesIO()
    writer.write(output)
    output.seek(0)
    
    return output.read()


def encrypt_word_msoffcrypto(input_docx_bytes: bytes, password: str) -> bytes:
    """
    使用msoffcrypto-tool加密Word文档
    
    Args:
        input_docx_bytes: 输入Word字节内容
        password: 加密密码
        
    Returns:
        加密后的Word字节内容
    """
    try:
        import msoffcrypto
    except ImportError:
        print("请安装msoffcrypto-tool: pip install msoffcrypto-tool")
        raise
    
    # 读取Word文档
    input_stream = io.BytesIO(input_docx_bytes)
    output_stream = io.BytesIO()
    
    # 创建加密对象
    office_file = msoffcrypto.OfficeFile(input_stream)
    
    # 设置密码加密
    # 注意：load_key是用于解密的，加密时直接调用encrypt即可
    office_file.encrypt(password, output_stream)
    
    output_stream.seek(0)
    return output_stream.read()


def create_simple_pdf() -> bytes:
    """创建一个简单的PDF用于测试"""
    from PyPDF2 import PdfWriter
    from io import BytesIO
    
    # 创建一个空白PDF
    writer = PdfWriter()
    
    # 添加一个空白页面
    writer.add_blank_page(width=612, height=792)  # Letter size
    
    output = BytesIO()
    writer.write(output)
    output.seek(0)
    return output.read()


def main():
    """演示文档加密功能"""
    print("=" * 60)
    print("文档加密MVP示例")
    print("=" * 60)
    
    output_dir = Path(__file__).parent / "outputs"
    output_dir.mkdir(exist_ok=True)
    
    # 1. 测试PDF加密
    print("\n1. 测试PDF加密...")
    try:
        # 创建一个简单的PDF
        pdf_bytes = create_simple_pdf()
        print(f"✓ 创建测试PDF: {len(pdf_bytes)} bytes")
        
        # 保存原始PDF
        original_pdf_path = output_dir / "original_document.pdf"
        with open(original_pdf_path, "wb") as f:
            f.write(pdf_bytes)
        print(f"  保存原始PDF到: {original_pdf_path}")
        
        # 加密PDF
        password = "test123"
        encrypted_pdf = encrypt_pdf_pypdf2(pdf_bytes, password)
        print(f"✓ PDF加密成功: {len(encrypted_pdf)} bytes")
        print(f"  密码: {password}")
        
        # 保存加密的PDF
        encrypted_pdf_path = output_dir / "encrypted_document.pdf"
        with open(encrypted_pdf_path, "wb") as f:
            f.write(encrypted_pdf)
        print(f"  保存加密PDF到: {encrypted_pdf_path}")
        
    except Exception as e:
        print(f"✗ PDF加密失败: {e}")
        import traceback
        traceback.print_exc()
    
    # 2. 测试Word加密
    print("\n2. 测试Word加密...")
    try:
        from docx import Document as DocxDocument
        
        doc = DocxDocument()
        doc.add_heading('这是一个需要加密的Word文档', 0)
        doc.add_paragraph('这个文档将被密码保护。')
        doc.add_paragraph('只有输入正确密码才能打开。')
        
        # 保存到字节流
        docx_stream = io.BytesIO()
        doc.save(docx_stream)
        docx_stream.seek(0)
        docx_bytes = docx_stream.read()
        
        print(f"✓ 创建测试Word文档: {len(docx_bytes)} bytes")
        
        # 保存原始Word文档
        original_docx_path = output_dir / "original_document.docx"
        with open(original_docx_path, "wb") as f:
            f.write(docx_bytes)
        print(f"  保存原始Word到: {original_docx_path}")
        
        # 加密Word文档
        password = "test456"
        encrypted_docx = encrypt_word_msoffcrypto(docx_bytes, password)
        print(f"✓ Word加密成功: {len(encrypted_docx)} bytes")
        print(f"  密码: {password}")
        
        # 保存加密的Word文档
        encrypted_docx_path = output_dir / "encrypted_document.docx"
        with open(encrypted_docx_path, "wb") as f:
            f.write(encrypted_docx)
        print(f"  保存加密Word到: {encrypted_docx_path}")
        
    except Exception as e:
        print(f"✗ Word加密失败: {e}")
        import traceback
        traceback.print_exc()
    
    # 3. 使用说明
    print("\n" + "=" * 60)
    print("验证加密效果:")
    print("=" * 60)
    print("1. 打开 mvp/outputs/original_document.pdf - 可以直接打开")
    print("2. 打开 mvp/outputs/encrypted_document.pdf - 需要输入密码: test123")
    print("3. 打开 mvp/outputs/original_document.docx - 可以直接打开")
    print("4. 打开 mvp/outputs/encrypted_document.docx - 需要输入密码: test456")
    
    print("\n" + "=" * 60)
    print("测试完成!")
    print("=" * 60)


if __name__ == "__main__":
    main()

